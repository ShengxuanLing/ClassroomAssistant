"""Task 21 fixture generator (deterministic, offline).

Generates all PDF/DOCX fixtures under tests/fixtures/documents/.
Run:  python scripts/gen_task21_fixtures.py
"""
import io
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
OUT = ROOT / "tests" / "fixtures" / "documents"
OUT.mkdir(parents=True, exist_ok=True)

from pypdf import PdfReader, PdfWriter
from pypdf.generic import ArrayObject, DecodedStreamObject, DictionaryObject, NameObject
from docx import Document as DocxDocument


def make_page(w, text_lines):
    p = w.add_blank_page(width=612, height=792)
    font = w._add_object(DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica")}))
    resources = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    p[NameObject("/Resources")] = w._add_object(resources)
    parts = []
    if text_lines:
        body = "BT /F1 11 Tf 72 720 Td 14 TL\n"
        for ln in text_lines:
            ln = ln.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            body += "(%s) Tj T*\n" % ln
        body += "ET\n"
        s = DecodedStreamObject()
        s.set_data(body.encode("latin-1"))
        parts.append(w._add_object(s))
    p[NameObject("/Contents")] = w._add_object(ArrayObject(parts))
    return p


def save_pdf(name, pages_texts, encrypt=None):
    w = PdfWriter()
    for text in pages_texts:
        lines = text.split("\n") if text else []
        make_page(w, lines)
    buf = io.BytesIO()
    w.write(buf)
    raw = buf.getvalue()
    if encrypt:
        reader = PdfReader(io.BytesIO(raw))
        out = PdfWriter()
        for pg in reader.pages:
            out.add_page(pg)
        out.encrypt(encrypt)
        buf2 = io.BytesIO()
        out.write(buf2)
        raw = buf2.getvalue()
    (OUT / name).write_bytes(raw)
    print(name, len(raw))


save_pdf("simple.pdf", [
    "Programa de la asignatura: Calculo I\nUnidad 1: Limites y continuidad. Ejemplo: f(x) = 1/x",
    "Fecha del examen: 20 de diciembre. Nota minima: 4.0",
])
save_pdf("multilingual.pdf", [
    "Introduccio al tema (ca). Definicion de funcion (es).",
    "Preguntas de revision: 1) Que es un limite? 2) Cuando existe?",
])
save_pdf("empty.pdf", [""])
save_pdf("scanned_like.pdf", ["", ""])
(OUT / "corrupted.pdf").write_bytes(b"%PDF-1.4 garbage not a real pdf body")
save_pdf("password.pdf", ["contenido protegido con password"], encrypt="secret")
(OUT / "zero_bytes.pdf").write_bytes(b"")
(OUT / "zero_bytes.docx").write_bytes(b"")
save_pdf("large_120_pages.pdf", [
    "Pagina %d: contenido de prueba de la leccion %d" % (i, i) for i in range(1, 121)
])


def save_docx(name, build):
    d = DocxDocument()
    build(d)
    buf = io.BytesIO()
    d.save(buf)
    (OUT / name).write_bytes(buf.getvalue())
    print(name, buf.tell())


def build_simple(d):
    d.add_paragraph("Titulo del documento: Calculus I")
    d.add_paragraph("Primera leccion sobre limites y continuidad.")
    d.add_paragraph("")
    d.add_paragraph("Segunda leccion: derivadas basicas.")


save_docx("simple.docx", build_simple)


def build_multilingual(d):
    d.add_paragraph("Introduccio al tema (ca)")
    d.add_paragraph("Definicion: funcion f: A -> B (es)")
    d.add_paragraph("Exemple: f(x) = x^2 + 1")


save_docx("multilingual.docx", build_multilingual)


def build_headings(d):
    d.add_heading("Capitulo 1: Analisis", level=1)
    d.add_paragraph("Texto de apoyo del capitulo.")
    d.add_heading("1.1 Limites", level=2)
    d.add_paragraph("Definicion eps-delta explicada.")
    d.add_heading("Capitulo 2: Derivadas", level=1)
    d.add_paragraph("Reglas de derivacion: suma, producto, cadena.")


save_docx("headings.docx", build_headings)


def build_tables(d):
    d.add_paragraph("Tabla de evaluacion")
    t = d.add_table(rows=4, cols=3)
    for i, row in enumerate([
        ["Asignacion", "Peso", "Fecha"],
        ["Examen 1", "30%", "Noviembre"],
        ["Trabajo", "40%", "Diciembre"],
        ["Final", "30%", "Enero"],
    ]):
        for j, val in enumerate(row):
            t.cell(i, j).text = val
    d.add_paragraph("Fin de la tabla.")
    t2 = d.add_table(rows=2, cols=2)
    t2.cell(0, 0).text = "A"
    t2.cell(0, 1).text = "B"
    t2.cell(1, 0).text = "C"
    t2.cell(1, 1).text = "D"


save_docx("tables.docx", build_tables)


def build_empty(d):
    pass


save_docx("empty.docx", build_empty)


def build_large_1000(d):
    for i in range(1000):
        d.add_paragraph("Linea %d de contenido repetido para proba de volumen." % i)


save_docx("large_1000.docx", build_large_1000)


def build_images(d):
    d.add_paragraph("Documento con imagenes insertadas.")
    png_path = Path(__file__).parent / "_tiny.png"
    d.add_picture(str(png_path))
    d.add_paragraph("Segunda imagen despues del texto.")
    d.add_picture(str(png_path))


save_docx("images.docx", build_images)

(OUT / "corrupted.docx").write_bytes(b"not a zip file, just junk bytes for corruption tests")
(OUT / "unsupported.txt").write_text("no", encoding="utf-8")
(OUT / "unsupported.py").write_text("print(1)", encoding="utf-8")

print("fixtures done:", len(os.listdir(OUT)))
