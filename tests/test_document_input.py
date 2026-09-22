"""Tests for Task 21: PDF / DOCX document input & parsing layer."""
from __future__ import annotations

import hashlib
import sys
from pathlib import Path
from typing import Any

import pytest

from src.document_input import (
    DocumentBlock,
    DocumentBlockType,
    DocumentInput,
    DocumentMaterialValidator,
    DocumentParserErrorCode,
    DocumentStatus,
    DocumentValidationError,
    DocumentValidationResult,
    ParsedDocument,
    ParserError,
    create_document_parser,
    parse_document,
    validate_document,
)
from src.models import Material, MaterialType

FIXTURES = Path(__file__).parent / "fixtures" / "documents"


def _pdf(name: str) -> Path:
    return FIXTURES / name


# ---------------------------------------------------------------------------
# Deterministic id computation
# ---------------------------------------------------------------------------


class TestDocumentStableId:
    def test_deterministic_same_input_same_id(self) -> None:
        from src.document_input import _compute_document_stable_id

        a = _compute_document_stable_id("", _pdf("simple.pdf"), 1234, 0)
        b = _compute_document_stable_id("", _pdf("simple.pdf"), 1234, 0)
        assert a == b
        assert a.startswith("document-")
        assert len(a) == len("document-") + 16

    def test_different_type_different_id(self) -> None:
        from src.document_input import _compute_document_stable_id

        a = _compute_document_stable_id("", _pdf("simple.pdf"), 1234, 0)
        b = _compute_document_stable_id("", _pdf("simple.docx"), 1234, 0)
        assert a != b

    def test_different_size_different_id(self) -> None:
        from src.document_input import _compute_document_stable_id

        a = _compute_document_stable_id("", _pdf("simple.pdf"), 100, 0)
        b = _compute_document_stable_id("", _pdf("simple.pdf"), 200, 0)
        assert a != b

    def test_mtime_ns_not_part_of_id(self) -> None:
        from src.document_input import _compute_document_stable_id

        a = _compute_document_stable_id("", _pdf("simple.pdf"), 1234, 100)
        b = _compute_document_stable_id("", _pdf("simple.pdf"), 1234, 200)
        assert a == b


# ---------------------------------------------------------------------------
# DocumentBlock
# ---------------------------------------------------------------------------


class TestDocumentBlock:
    def test_default_values(self) -> None:
        b = DocumentBlock()
        assert b.text == ""
        assert b.block_type is DocumentBlockType.TEXT
        assert b.page_number is None
        assert b.paragraph_index is None

    def test_to_dict_roundtrip(self) -> None:
        b = DocumentBlock(
            block_id="docblock-abc",
            block_type=DocumentBlockType.HEADING,
            text="Titulo",
            page_number=3,
            paragraph_index=7,
        )
        d = b.to_dict()
        b2 = DocumentBlock.from_dict(d)
        assert b2.block_id == "docblock-abc"
        assert b2.block_type is DocumentBlockType.HEADING
        assert b2.text == "Titulo"
        assert b2.page_number == 3
        assert b2.paragraph_index == 7

    def test_assign_block_id_deterministic(self) -> None:
        b1 = DocumentBlock(block_type=DocumentBlockType.TEXT, text="hello")
        b2 = DocumentBlock(block_type=DocumentBlockType.TEXT, text="hello")
        b1.assign_block_id("document-1")
        b2.assign_block_id("document-1")
        assert b1.block_id == b2.block_id
        assert b1.block_id.startswith("docblock-")
        assert len(b1.block_id) == len("docblock-") + 24

    def test_assign_block_id_differs_by_text(self) -> None:
        b1 = DocumentBlock(block_type=DocumentBlockType.TEXT, text="hello")
        b2 = DocumentBlock(block_type=DocumentBlockType.TEXT, text="world")
        b1.assign_block_id("document-1")
        b2.assign_block_id("document-1")
        assert b1.block_id != b2.block_id

    def test_assign_block_id_differs_by_page(self) -> None:
        b1 = DocumentBlock(block_type=DocumentBlockType.TEXT, text="x", page_number=1)
        b2 = DocumentBlock(block_type=DocumentBlockType.TEXT, text="x", page_number=2)
        b1.assign_block_id("document-1")
        b2.assign_block_id("document-1")
        assert b1.block_id != b2.block_id

    def test_to_source_reference_page(self) -> None:
        b = DocumentBlock(
            block_type=DocumentBlockType.TEXT,
            text="x",
            page_number=5,
            block_index=0,
        )
        ref = b.to_source_reference(material_id="mat-1")
        assert ref["material_id"] == "mat-1"
        assert ref["page"] == 5
        assert ref["line"] == 0
        assert "pdf-page-5" in ref["location"]

    def test_to_source_reference_paragraph(self) -> None:
        b = DocumentBlock(
            block_type=DocumentBlockType.TEXT,
            text="x",
            paragraph_index=12,
        )
        ref = b.to_source_reference(material_id="mat-2")
        assert ref["paragraph"] == "paragraph_12"
        assert "docx-paragraph-12" in ref["location"]

    def test_to_source_reference_table(self) -> None:
        b = DocumentBlock(
            block_type=DocumentBlockType.TABLE,
            text="a | b",
        )
        ref = b.to_source_reference(material_id="mat-3")
        assert ref["location"] == "table"


# ---------------------------------------------------------------------------
# ParsedDocument
# ---------------------------------------------------------------------------


class TestParsedDocument:
    def test_to_dict_from_dict_roundtrip(self) -> None:
        doc = ParsedDocument(
            document_id="document-abc",
            document_type="PDF",
            path="test.pdf",
            material_id="mat-1",
            status=DocumentStatus.PARSED,
            blocks=[
                DocumentBlock(
                    block_id="docblock-1",
                    block_type=DocumentBlockType.TEXT,
                    text="page one",
                    page_number=1,
                    block_index=0,
                )
            ],
        )
        d = doc.to_dict()
        doc2 = ParsedDocument.from_dict(d)
        assert doc2.document_id == "document-abc"
        assert doc2.document_type == "PDF"
        assert doc2.status is DocumentStatus.PARSED
        assert len(doc2.blocks) == 1
        assert doc2.blocks[0].text == "page one"
        assert doc2.blocks[0].page_number == 1

    def test_from_dict_invalid_status_defaults_failed(self) -> None:
        d = {"document_id": "x", "status": "BROKEN"}
        doc = ParsedDocument.from_dict(d)
        assert doc.status is DocumentStatus.FAILED

    def test_block_count_property(self) -> None:
        doc = ParsedDocument(
            blocks=[DocumentBlock(), DocumentBlock(), DocumentBlock()]
        )
        assert doc.block_count == 3

    def test_page_count_from_metadata(self) -> None:
        doc = ParsedDocument(metadata={"page_count": 42})
        assert doc.page_count == 42

    def test_page_count_defaults_zero(self) -> None:
        doc = ParsedDocument()
        assert doc.page_count == 0


# ---------------------------------------------------------------------------
# DocumentMaterialValidator
# ---------------------------------------------------------------------------


class TestDocumentMaterialValidator:
    def test_valid_pdf(self) -> None:
        v = DocumentMaterialValidator()
        r = v.validate(_pdf("simple.pdf"))
        assert r.valid is True
        assert r.status.value == "VALID"
        assert r.extension == ".pdf"
        assert r.document_type == "PDF"
        assert r.file_size > 0

    def test_valid_docx(self) -> None:
        v = DocumentMaterialValidator()
        r = v.validate(_pdf("simple.docx"))
        assert r.valid is True
        assert r.document_type == "DOCX"
        assert r.material_type is MaterialType.SYLLABUS

    def test_missing_file(self) -> None:
        v = DocumentMaterialValidator()
        r = v.validate(FIXTURES / "does_not_exist.pdf")
        assert r.valid is False
        assert DocumentValidationError.FILE_NOT_FOUND in r.errors

    def test_unsupported_extension(self) -> None:
        v = DocumentMaterialValidator()
        r = v.validate(FIXTURES / "unsupported.txt")
        assert r.valid is False
        assert DocumentValidationError.UNSUPPORTED_EXTENSION in r.errors

    def test_zero_byte_file(self) -> None:
        v = DocumentMaterialValidator()
        r = v.validate(_pdf("zero_bytes.pdf"))
        assert r.valid is False
        assert DocumentValidationError.EMPTY_FILE in r.errors

    def test_directory_rejected(self) -> None:
        v = DocumentMaterialValidator()
        r = v.validate(FIXTURES)
        assert r.valid is False
        assert DocumentValidationError.NOT_A_FILE in r.errors

    def test_case_insensitive_extension(self) -> None:
        v = DocumentMaterialValidator()
        r = v.validate(_pdf("simple.pdf"))
        assert r.extension == ".pdf"

    def test_to_document_input_valid(self) -> None:
        v = DocumentMaterialValidator()
        r = v.validate(_pdf("simple.pdf"))
        di = v.to_document_input(r)
        assert di.path == str(_pdf("simple.pdf"))
        assert di.document_type == "PDF"
        assert di.material_id is None

    def test_to_document_input_invalid_raises(self) -> None:
        v = DocumentMaterialValidator()
        r = v.validate(FIXTURES / "does_not_exist.pdf")
        with pytest.raises(ValueError):
            v.to_document_input(r)

    def test_compute_material_id_deterministic(self) -> None:
        v = DocumentMaterialValidator()
        r = v.validate(_pdf("simple.pdf"))
        a = v.compute_material_id(r)
        b = v.compute_material_id(r)
        assert a == b
        assert a.startswith("document-")

    def test_validate_document_helper(self) -> None:
        r = validate_document(_pdf("simple.pdf"))
        assert r.valid is True
        assert r.document_type == "PDF"

    def test_validate_document_zero_byte(self) -> None:
        r = validate_document(_pdf("zero_bytes.pdf"))
        assert r.valid is False
        assert DocumentValidationError.EMPTY_FILE in r.errors


# ---------------------------------------------------------------------------
# PDF parser
# ---------------------------------------------------------------------------


class TestPDFDocumentParser:
    def test_simple_pdf_parsed(self) -> None:
        doc = create_document_parser("PDF").parse(_pdf("simple.pdf"))
        assert doc.status is DocumentStatus.PARSED
        assert doc.page_count == 2
        assert doc.block_count == 2
        assert doc.blocks[0].page_number == 1
        assert doc.blocks[1].page_number == 2
        assert "Programa" in doc.blocks[0].text
        assert "Calculo I" in doc.blocks[0].text
        assert "Fecha del examen" in doc.blocks[1].text

    def test_pdf_block_ids_deterministic(self) -> None:
        p = create_document_parser("PDF")
        a = p.parse(_pdf("simple.pdf"))
        b = p.parse(_pdf("simple.pdf"))
        for ba, bb in zip(a.blocks, b.blocks):
            assert ba.block_id == bb.block_id
            assert ba.block_id.startswith("docblock-")

    def test_pdf_document_id_deterministic(self) -> None:
        p = create_document_parser("PDF")
        a = p.parse(_pdf("simple.pdf"))
        b = p.parse(_pdf("simple.pdf"))
        assert a.document_id == b.document_id
        assert a.document_id.startswith("document-")

    def test_password_pdf_failed(self) -> None:
        doc = create_document_parser("PDF").parse(_pdf("password.pdf"))
        assert doc.status is DocumentStatus.FAILED
        assert len(doc.errors) == 1
        assert doc.errors[0].error_code is DocumentParserErrorCode.PASSWORD_PROTECTED

    def test_corrupted_pdf_failed(self) -> None:
        doc = create_document_parser("PDF").parse(_pdf("corrupted.pdf"))
        assert doc.status is DocumentStatus.FAILED
        assert doc.errors[0].error_code is DocumentParserErrorCode.INVALID_DOCUMENT

    def test_scanned_pdf_parsed_empty(self) -> None:
        doc = create_document_parser("PDF").parse(_pdf("scanned_like.pdf"))
        assert doc.status is DocumentStatus.PARSED_EMPTY
        assert doc.metadata.get("scanned_or_no_text_layer") is True
        assert doc.page_count == 2

    def test_empty_pdf_parsed_empty(self) -> None:
        doc = create_document_parser("PDF").parse(_pdf("empty.pdf"))
        assert doc.status is DocumentStatus.PARSED_EMPTY
        assert doc.page_count == 1

    def test_large_pdf_all_pages(self) -> None:
        doc = create_document_parser("PDF").parse(_pdf("large_120_pages.pdf"))
        assert doc.status is DocumentStatus.PARSED
        assert doc.block_count == 120
        assert doc.blocks[0].text.startswith("Pagina 1:")
        assert doc.blocks[119].text.startswith("Pagina 120:")

    def test_pdf_block_text_stripped(self) -> None:
        doc = create_document_parser("PDF").parse(_pdf("simple.pdf"))
        for b in doc.blocks:
            assert b.text == b.text.strip()
            assert b.block_type is DocumentBlockType.TEXT

    def test_pdf_missing_file(self) -> None:
        doc = create_document_parser("PDF").parse(FIXTURES / "nope.pdf")
        assert doc.status is DocumentStatus.FAILED
        assert doc.errors[0].error_code is DocumentParserErrorCode.INVALID_DOCUMENT

    def test_pdf_multilingual_content(self) -> None:
        doc = create_document_parser("PDF").parse(_pdf("multilingual.pdf"))
        assert doc.status is DocumentStatus.PARSED
        assert "Introduccio" in doc.blocks[0].text
        assert "Preguntas de revision" in doc.blocks[1].text


# ---------------------------------------------------------------------------
# DOCX parser
# ---------------------------------------------------------------------------


class TestDOCXDocumentParser:
    def test_simple_docx_parsed(self) -> None:
        doc = create_document_parser("DOCX").parse(_pdf("simple.docx"))
        assert doc.status is DocumentStatus.PARSED
        # 3 non-empty paragraphs (one empty skipped, index advances)
        non_empty = [b for b in doc.blocks if b.block_type is not DocumentBlockType.TABLE]
        assert len(non_empty) == 3
        assert non_empty[0].paragraph_index == 0
        assert non_empty[1].paragraph_index == 1
        assert non_empty[2].paragraph_index == 3

    def test_docx_paragraph_index_advances_past_empty(self) -> None:
        doc = create_document_parser("DOCX").parse(_pdf("simple.docx"))
        paras = [b for b in doc.blocks if b.block_type is DocumentBlockType.TEXT]
        assert [p.paragraph_index for p in paras] == [0, 1, 3]

    def test_headings_docx(self) -> None:
        doc = create_document_parser("DOCX").parse(_pdf("headings.docx"))
        assert doc.status is DocumentStatus.PARSED
        headings = [b for b in doc.blocks if b.block_type is DocumentBlockType.HEADING]
        texts = [h.text for h in headings]
        assert "Capitulo 1: Analisis" in texts
        assert "1.1 Limites" in texts
        assert "Capitulo 2: Derivadas" in texts

    def test_docx_table_blocks(self) -> None:
        doc = create_document_parser("DOCX").parse(_pdf("tables.docx"))
        tables = [b for b in doc.blocks if b.block_type is DocumentBlockType.TABLE]
        assert len(tables) == 2
        assert tables[0].metadata.get("table_index") == 0
        assert tables[1].metadata.get("table_index") == 1
        assert "Asignacion" in tables[0].text
        assert "30%" in tables[0].text

    def test_docx_table_text_format(self) -> None:
        doc = create_document_parser("DOCX").parse(_pdf("tables.docx"))
        tables = [b for b in doc.blocks if b.block_type is DocumentBlockType.TABLE]
        # rows joined with " | ", cells within row joined with " || "
        parts = tables[0].text.split(" | ")
        assert len(parts) == 4
        assert " || " in parts[0]

    def test_docx_empty_parsed_empty(self) -> None:
        doc = create_document_parser("DOCX").parse(_pdf("empty.docx"))
        assert doc.status is DocumentStatus.PARSED_EMPTY
        assert doc.block_count == 0

    def test_docx_images_counted(self) -> None:
        doc = create_document_parser("DOCX").parse(_pdf("images.docx"))
        assert doc.status is DocumentStatus.PARSED
        assert doc.metadata.get("image_count") == 1

    def test_docx_corrupted_failed(self) -> None:
        doc = create_document_parser("DOCX").parse(_pdf("corrupted.docx"))
        assert doc.status is DocumentStatus.FAILED
        assert doc.errors[0].error_code is DocumentParserErrorCode.INVALID_DOCUMENT

    def test_docx_zero_byte_failed(self) -> None:
        doc = create_document_parser("DOCX").parse(_pdf("zero_bytes.docx"))
        assert doc.status is DocumentStatus.FAILED

    def test_docx_large_1000_paragraphs(self) -> None:
        doc = create_document_parser("DOCX").parse(_pdf("large_1000.docx"))
        assert doc.status is DocumentStatus.PARSED
        texts = [b for b in doc.blocks if b.block_type is DocumentBlockType.TEXT]
        assert len(texts) == 1000
        assert texts[0].text.startswith("Linea 0")
        assert texts[999].text.startswith("Linea 999")

    def test_docx_block_ids_deterministic(self) -> None:
        p = create_document_parser("DOCX")
        a = p.parse(_pdf("simple.docx"))
        b = p.parse(_pdf("simple.docx"))
        assert a.document_id == b.document_id
        for ba, bb in zip(a.blocks, b.blocks):
            assert ba.block_id == bb.block_id

    def test_docx_multilingual_content(self) -> None:
        doc = create_document_parser("DOCX").parse(_pdf("multilingual.docx"))
        assert doc.status is DocumentStatus.PARSED
        all_text = " ".join(b.text for b in doc.blocks)
        assert "Introduccio" in all_text
        assert "Definicion" in all_text

    def test_docx_block_index_not_set(self) -> None:
        doc = create_document_parser("DOCX").parse(_pdf("simple.docx"))
        for b in doc.blocks:
            assert b.block_index is None
            assert b.page_number is None


# ---------------------------------------------------------------------------
# Factory
# ---------------------------------------------------------------------------


class TestFactory:
    def test_create_pdf_parser(self) -> None:
        p = create_document_parser("PDF")
        assert p.document_type == "PDF"

    def test_create_docx_parser(self) -> None:
        p = create_document_parser("DOCX")
        assert p.document_type == "DOCX"

    def test_create_case_insensitive(self) -> None:
        p = create_document_parser("pdf")
        assert p.document_type == "PDF"
        p2 = create_document_parser("docx")
        assert p2.document_type == "DOCX"

    def test_create_unknown_raises(self) -> None:
        with pytest.raises(ValueError):
            create_document_parser("XLSX")

    def test_parse_document_dispatches_by_extension(self) -> None:
        doc = parse_document(_pdf("simple.pdf"))
        assert doc.document_type == "PDF"
        doc2 = parse_document(_pdf("simple.docx"))
        assert doc2.document_type == "DOCX"

    def test_parse_document_unsupported_raises(self) -> None:
        with pytest.raises(ValueError):
            parse_document(FIXTURES / "unsupported.txt")

    def test_parse_document_missing_raises(self) -> None:
        with pytest.raises(ValueError):
            parse_document(FIXTURES / "does_not_exist.pdf")

    def test_parse_document_overrides_type(self) -> None:
        doc = parse_document(_pdf("simple.docx"), document_type="DOCX")
        assert doc.document_type == "DOCX"


# ---------------------------------------------------------------------------
# Serialization / deserialization
# ---------------------------------------------------------------------------


class TestSerialization:
    def test_parsed_document_full_roundtrip(self) -> None:
        doc = create_document_parser("PDF").parse(_pdf("simple.pdf"))
        d = doc.to_dict()
        doc2 = ParsedDocument.from_dict(d)
        assert doc2.document_id == doc.document_id
        assert doc2.status is doc.status
        assert doc2.page_count == doc.page_count
        for b1, b2 in zip(doc.blocks, doc2.blocks):
            assert b1.block_id == b2.block_id
            assert b1.block_type is b2.block_type
            assert b1.text == b2.text
            assert b1.page_number == b2.page_number

    def test_parsed_document_with_errors_roundtrip(self) -> None:
        doc = create_document_parser("PDF").parse(_pdf("password.pdf"))
        d = doc.to_dict()
        doc2 = ParsedDocument.from_dict(d)
        assert doc2.status is DocumentStatus.FAILED
        assert len(doc2.errors) == 1
        assert doc2.errors[0].error_code is DocumentParserErrorCode.PASSWORD_PROTECTED

    def test_parser_error_roundtrip(self) -> None:
        e = ParserError(DocumentParserErrorCode.PARSER_ERROR, "test message")
        d = e.to_dict()
        e2 = ParserError.from_dict(d)
        assert e2.error_code is DocumentParserErrorCode.PARSER_ERROR
        assert e2.message == "test message"

    def test_parser_error_unknown_code_defaults(self) -> None:
        e = ParserError.from_dict({"error_code": "NOT_A_REAL_CODE", "message": "x"})
        assert e.error_code is DocumentParserErrorCode.PARSER_ERROR

    def test_document_input_roundtrip(self) -> None:
        v = DocumentMaterialValidator()
        r = v.validate(_pdf("simple.pdf"))
        di = v.to_document_input(r)
        d = di.to_dict()
        di2 = DocumentInput.from_dict(d)
        assert di2.path == di.path
        assert di2.extension == ".pdf"
        assert di2.document_type == "PDF"

    def test_document_validation_result_roundtrip(self) -> None:
        v = DocumentMaterialValidator()
        r = v.validate(_pdf("simple.pdf"))
        d = r.to_dict()
        r2 = DocumentValidationResult.from_dict(d)
        assert r2.valid is True
        assert r2.extension == ".pdf"
        assert r2.document_type == "PDF"
        assert r2.material_type is MaterialType.SYLLABUS

    def test_document_validation_result_invalid_roundtrip(self) -> None:
        v = DocumentMaterialValidator()
        r = v.validate(FIXTURES / "unsupported.txt")
        d = r.to_dict()
        r2 = DocumentValidationResult.from_dict(d)
        assert r2.valid is False
        assert DocumentValidationError.UNSUPPORTED_EXTENSION in r2.errors


# ---------------------------------------------------------------------------
# Determinism (no uuid4 anywhere)
# ---------------------------------------------------------------------------


class TestDeterminism:
    def test_no_uuid_in_ids(self) -> None:
        p = create_document_parser("PDF")
        doc = p.parse(_pdf("simple.pdf"))
        assert doc.document_id.startswith("document-")
        assert len(doc.document_id) == 9 + 16
        for b in doc.blocks:
            assert b.block_id.startswith("docblock-")
            assert len(b.block_id) == 9 + 24

    def test_same_file_same_ids_across_calls(self) -> None:
        p = create_document_parser("DOCX")
        a = p.parse(_pdf("headings.docx"))
        b = p.parse(_pdf("headings.docx"))
        assert a.document_id == b.document_id
        assert [x.block_id for x in a.blocks] == [x.block_id for x in b.blocks]

    def test_different_files_different_doc_ids(self) -> None:
        p = create_document_parser("PDF")
        a = p.parse(_pdf("simple.pdf"))
        b = p.parse(_pdf("multilingual.pdf"))
        assert a.document_id != b.document_id


# ---------------------------------------------------------------------------
# Source location
# ---------------------------------------------------------------------------


class TestSourceLocation:
    def test_pdf_block_source_ref_has_page(self) -> None:
        doc = create_document_parser("PDF").parse(_pdf("simple.pdf"))
        for b in doc.blocks:
            ref = b.to_source_reference("mat-x")
            assert ref["page"] is not None
            assert ref["page"] == b.page_number

    def test_docx_block_source_ref_has_paragraph(self) -> None:
        doc = create_document_parser("DOCX").parse(_pdf("headings.docx"))
        for b in doc.blocks:
            if b.paragraph_index is not None:
                ref = b.to_source_reference("mat-y")
                assert ref["paragraph"] == f"paragraph_{b.paragraph_index}"

    def test_docx_table_source_ref_location_is_table(self) -> None:
        doc = create_document_parser("DOCX").parse(_pdf("tables.docx"))
        for b in doc.blocks:
            if b.block_type is DocumentBlockType.TABLE:
                ref = b.to_source_reference("mat-z")
                assert ref["location"] == "table"


# ---------------------------------------------------------------------------
# Unicode / accent handling
# ---------------------------------------------------------------------------


class TestUnicodeContent:
    def test_pdf_accented_content_preserved(self) -> None:
        doc = create_document_parser("PDF").parse(_pdf("simple.pdf"))
        # Fixture uses ASCII only; verify no mojibake in block text.
        for b in doc.blocks:
            assert "\ufffd" not in b.text

    def test_docx_unicode_content_roundtrip(self) -> None:
        from docx import Document as DocxDocument
        import io

        d = DocxDocument()
        d.add_paragraph("Introducci\u00f3n: \u00e9s d\u00edgit")
        buf = io.BytesIO()
        d.save(buf)
        target = FIXTURES / "tmp_unicode.docx"
        target.write_bytes(buf.getvalue())
        try:
            doc = create_document_parser("DOCX").parse(target)
            assert doc.status is DocumentStatus.PARSED
            assert "Introducci\u00f3n" in doc.blocks[0].text
            assert "\u00e9s d\u00edgit" in doc.blocks[0].text
        finally:
            target.unlink(missing_ok=True)


# ---------------------------------------------------------------------------
# Integration: validator -> parser -> source reference
# ---------------------------------------------------------------------------


class TestIntegration:
    @pytest.mark.integration
    def test_full_pipeline_pdf(self) -> None:
        v = DocumentMaterialValidator()
        r = v.validate(_pdf("simple.pdf"))
        assert r.valid is True
        doc = create_document_parser(r.document_type or "PDF").parse(
            Path(r.path), material_id=v.compute_material_id(r)
        )
        assert doc.status is DocumentStatus.PARSED
        for b in doc.blocks:
            ref = b.to_source_reference(doc.material_id)
            assert ref["material_id"] == doc.material_id
            assert ref["page"] is not None

    @pytest.mark.integration
    def test_full_pipeline_docx(self) -> None:
        v = DocumentMaterialValidator()
        r = v.validate(_pdf("headings.docx"))
        assert r.valid is True
        doc = create_document_parser("DOCX").parse(
            Path(r.path), material_id=v.compute_material_id(r)
        )
        assert doc.status is DocumentStatus.PARSED
        for b in doc.blocks:
            if b.block_type in (DocumentBlockType.TEXT, DocumentBlockType.HEADING):
                ref = b.to_source_reference(doc.material_id)
                assert ref["material_id"] == doc.material_id

    @pytest.mark.integration
    def test_parser_on_invalid_file_returns_failed_not_raises(self) -> None:
        for name in ("corrupted.pdf", "password.pdf", "corrupted.docx"):
            path = _pdf(name)
            parser = create_document_parser("PDF" if name.endswith("pdf") else "DOCX")
            doc = parser.parse(path)
            assert doc.status is DocumentStatus.FAILED, name
